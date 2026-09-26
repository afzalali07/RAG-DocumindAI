"""One logical page per sheet; preserve cell positions and report truncation."""
from pathlib import Path

from app.parsers.base import ExtractedTable, PageSegment, ParseError

_MAX_ROWS_PER_SHEET = 2000
_MAX_COLUMNS = 100


def parse(path: Path) -> list[PageSegment]:
    from openpyxl import load_workbook

    segments = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for number, sheet in enumerate(wb.worksheets, 1):
                # Exported workbooks can omit dimensions or incorrectly declare A1.
                # Read actual XML rows with explicit safety bounds instead.
                sheet.reset_dimensions()
                rows = [
                    ["" if cell is None else str(cell) for cell in row]
                    for row in sheet.iter_rows(
                        max_row=_MAX_ROWS_PER_SHEET + 1,
                        max_col=_MAX_COLUMNS + 1, values_only=True,
                    )
                ]
                truncated = any(any(cell for cell in row) for row in rows[_MAX_ROWS_PER_SHEET:]) or any(
                    any(row[_MAX_COLUMNS:]) for row in rows
                )
                rows = [row[:_MAX_COLUMNS] for row in rows[:_MAX_ROWS_PER_SHEET]]
                while rows and not any(rows[-1]):
                    rows.pop()
                width = max((i + 1 for row in rows for i, cell in enumerate(row) if cell), default=0)
                rows = [row[:width] for row in rows]
                tables = [ExtractedTable(rows, sheet.title)] if width else []
                warnings = []
                if truncated:
                    warnings.append("Extraction limited to the first 2,000 rows and 100 columns.")
                segments.append(PageSegment(number, tables[0].as_text() if tables else "", sheet.title, tables, warnings))
        finally:
            wb.close()
    except Exception as exc:
        raise ParseError(f"Could not read Excel workbook: {exc}") from exc
    if not any(segment.text for segment in segments):
        raise ParseError("Excel workbook is empty.")
    return segments
