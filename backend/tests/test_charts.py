import json
import shutil
import pytest
from openpyxl import Workbook
from app.core.charts import ChartPlan, calculate
from app.core.spreadsheet import read_workbook

@pytest.fixture
def workbook(tmp_path):
    book = Workbook(); sheet = book.active; sheet.title = 'Data'
    sheet.append(['Factory', 'Emissions', 'CO5 (20)', 'CO5 (100)'])
    sheet.append(['A', 10, 2, 10]); sheet.append(['B', 30, 4, 20]); sheet.append(['A', 20, 6, 30]); sheet.append(['C', 0, 0, 0])
    path = tmp_path / 'data.xlsx'; book.save(path)
    return path

def test_group_average_and_row_references(workbook):
    result = calculate(read_workbook(workbook), ChartPlan(label_column='Factory',value_column='Emissions',aggregation='average'), 'bar')
    assert result['points'][0] == {'label':'A','value':15,'rows':[2,4]}
    assert result['points'][2]['value'] == 0

def test_filters_and_top_results(workbook):
    result = calculate(read_workbook(workbook), ChartPlan(label_column='Factory',value_column='CO5 (100)',filters=[{'column':'Emissions','operator':'ge','value':10}],order='descending',limit=1), 'line')
    assert result['points'][0]['value'] == 30
    assert result['total_points'] == 3

def test_count(workbook):
    result = calculate(read_workbook(workbook), ChartPlan(label_column='Factory',aggregation='count'), 'pie')
    assert result['points'][0]['value'] == 2

@pytest.mark.parametrize('value', [-1, 0])
def test_pie_rejects_invalid_values(workbook,value):
    sheets=read_workbook(workbook); sheets[0]['rows']=[(2,['A',value,0,0])]
    with pytest.raises(ValueError, match='Pie charts'):
        calculate(sheets, ChartPlan(label_column='Factory', value_column='Emissions'), 'pie')

def test_missing_formulas_and_ambiguous_headers(workbook):
    sheets=read_workbook(workbook); sheets[0]['missing'].add((2,1))
    with pytest.raises(ValueError,match='formulas'):
        calculate(sheets,ChartPlan(label_column='Factory',value_column='Emissions'),'bar')
    with pytest.raises(ValueError,match='exact column'):
        calculate(sheets,ChartPlan(label_column='Factory',value_column='CO5'),'bar')

def test_chart_api(client,workbook,monkeypatch):
    from app.api.routes import charts
    from app.config import get_settings
    from app.db.models import Document
    from app.db.session import SessionLocal
    class Planner:
        async def stream_json(self,*args):
            yield json.dumps({'label_column':'Factory','value_column':'Emissions','aggregation':'sum'})
    monkeypatch.setattr(charts,'get_provider',lambda _:Planner())
    with SessionLocal() as db:
        doc=Document(filename='data.xlsx',status='ready',content_type='application/octet-stream')
        db.add(doc); db.commit(); doc_id=doc.id
    shutil.copyfile(workbook,get_settings().upload_dir/f'{doc_id}.xlsx')
    response=client.post('/api/spreadsheet/visualize',json={'document_id':doc_id,'query':'Total emissions by factory','chart_type':'bar'})
    assert response.status_code==200
    assert response.json()['points'][0]['value']==30
    assert client.post('/api/spreadsheet/visualize',json={'document_id':'missing','query':'Plot data'}).status_code==404
    assert client.post('/api/spreadsheet/visualize',json={'document_id':doc_id,'query':'Plot data','chart_type':'invalid'}).status_code==422

def test_aggregate_query_uses_matching_sheet_without_model(workbook):
    import asyncio
    from app.core.charts import visualize
    result = asyncio.run(visualize(workbook, 'total emissions by factory', 'bar', None))
    assert result['points'][0]['value'] == 30
    assert result['aggregation'] == 'sum'


def test_ambiguous_emissions_lists_measures(tmp_path):
    import asyncio
    from app.core.charts import visualize
    book=Workbook(); sheet=book.active; sheet.title='Emission Data'
    sheet.append(['Factory','CO2 (t)','SO2 (kg)']); sheet.append(['A',10,200])
    book.create_sheet('Summary').append(['Dashboard'])
    path=tmp_path/'emissions.xlsx'; book.save(path)
    with pytest.raises(ValueError, match='Available measures: CO2'):
        asyncio.run(visualize(path,'total emissions by factory','bar',None))
    result=asyncio.run(visualize(path,'Total CO2 (t) by Factory','bar',None))
    assert result['sheet']=='Emission Data'
    assert result['points'][0]['value']==10


def test_echoed_clarification_is_actionable(workbook):
    import asyncio
    from app.core.charts import visualize
    class Echo:
        async def stream_json(self,*args):
            yield json.dumps({'clarification':'Plot something'})
    with pytest.raises(ValueError,match='Available worksheet columns'):
        asyncio.run(visualize(workbook,'Plot something','bar',Echo()))
