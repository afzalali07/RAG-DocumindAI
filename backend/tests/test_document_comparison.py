"""Extraction provenance, column alignment, and isolated two-document comparison."""
import io
import json

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle

from app.core import vectorstore
from app.core.extraction import cache_path
from app.db.models import Document
from app.db.session import SessionLocal
from app.parsers.base import page_count, parse_document


def pdf_with_table():
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(50, 800, "Contract overview")
    pdf.showPage()
    table = Table([["Item", "Quantity", "Price"], ["Wheat", "", "120"], ["Rice", "5", "80"]])
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 1, 'black')]))
    table.wrapOn(pdf, 500, 700)
    table.drawOn(pdf, 50, 650)
    pdf.showPage()
    pdf.showPage()  # Deliberate trailing blank page must remain counted.
    pdf.save()
    return output.getvalue()


def test_pdf_table_page_and_blank_page(tmp_path):
    path = tmp_path / 'table.pdf'
    path.write_bytes(pdf_with_table())
    pages = parse_document(path, 'application/pdf')
    assert page_count(pages) == 3
    assert not pages[0].tables
    assert pages[1].page == 2
    assert pages[1].tables[0].rows[1] == ['Wheat', '', '120']
    assert pages[2].text == '' and pages[2].warnings


def test_word_tables_keep_document_order(tmp_path):
    doc = WordDocument()
    doc.add_paragraph('Before table')
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Product', 'Qty', 'Price']):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ['Wheat', '', '100']):
        cell.text = text
    doc.add_paragraph('After table')
    path = tmp_path / 'table.docx'
    doc.save(path)
    pages = parse_document(path, 'application/octet-stream')
    assert pages[0].label == 'Section 1'
    assert pages[0].tables[0].rows[1] == ['Wheat', '', '100']
    assert pages[0].text.index('Before table') < pages[0].text.index('Wheat') < pages[0].text.index('After table')


def test_excel_preserves_blank_cells_and_sheets(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Prices'
    sheet.append(['Product', 'Qty', 'Price'])
    sheet.append(['Rice', None, 0])
    workbook.create_sheet('Empty')
    path = tmp_path / 'table.xlsx'
    workbook.save(path)
    pages = parse_document(path, 'application/octet-stream')
    assert len(pages) == 2 and pages[1].label == 'Empty'
    assert pages[0].tables[0].rows[1] == ['Rice', '', '0']


@pytest.fixture
def comparison_pair(client, monkeypatch):
    # Exercise upload/parsing/cache/DB/SSE with no external embedding or LLM calls.
    monkeypatch.setattr(vectorstore, 'add_document_chunks', lambda **kwargs: None)
    monkeypatch.setattr(vectorstore, 'delete_document', lambda *args: None)
    ids = []
    for name in ['contract-a.pdf', 'contract-b.pdf']:
        response = client.post('/api/documents', files={'file': (name, pdf_with_table(), 'application/pdf')})
        assert response.status_code == 201
        doc_id = response.json()['id']
        assert client.get(f'/api/documents/{doc_id}').json()['status'] == 'ready'
        ids.append(doc_id)
    yield ids
    for doc_id in ids:
        client.delete(f'/api/documents/{doc_id}')


def events(response):
    parsed = {}
    for frame in response.text.split('\n\n'):
        lines = frame.splitlines()
        if len(lines) >= 2:
            parsed[lines[0].removeprefix('event: ')] = json.loads(lines[1].removeprefix('data: '))
    return parsed


def test_extraction_existing_upload_and_delete(client, comparison_pair):
    doc_id = comparison_pair[0]
    cache_path(doc_id).unlink()  # Simulate a document uploaded before this feature.
    response = client.get(f'/api/documents/{doc_id}/extraction')
    assert response.status_code == 200
    data = response.json()
    assert data['page_count'] == 3 and data['table_count'] == 1
    assert data['pages'][1]['tables'][0]['rows'][1] == ['Wheat', '', '120']
    assert cache_path(doc_id).is_file()
    client.delete(f'/api/documents/{doc_id}')
    assert not cache_path(doc_id).exists()
    assert client.get(f'/api/documents/{doc_id}/extraction').status_code == 404


@pytest.mark.parametrize('ids', [[], ['a'], ['a', 'a'], ['a', 'b', 'c']])
def test_comparison_requires_two_distinct_documents(client, ids):
    response = client.post('/api/chat', json={'mode': 'compare', 'message': 'Compare', 'document_ids': ids})
    assert response.status_code == 422


def test_comparison_rejects_missing_document(client):
    response = client.post('/api/chat', json={'mode': 'compare', 'message': 'Compare', 'document_ids': ['a', 'b']})
    assert response.status_code == 404


def test_comparison_rejects_processing_document(client, comparison_pair):
    with SessionLocal() as db:
        doc = db.get(Document, comparison_pair[0])
        doc.status = 'processing'
        db.commit()
    response = client.post('/api/chat', json={'mode': 'compare', 'message': 'Compare', 'document_ids': comparison_pair})
    assert response.status_code == 409
    assert client.get(f'/api/documents/{comparison_pair[0]}/extraction').status_code == 409


def test_comparison_cites_both_documents_and_persists(client, comparison_pair):
    response = client.post('/api/chat', json={
        'mode': 'compare', 'message': 'Compare prices', 'document_ids': comparison_pair,
        'model': 'mock', 'lang': 'en', 'category': 'Unrelated category',
    })
    assert response.status_code == 200
    result = events(response)
    assert 'error' not in result
    sources = result['sources']['sources']
    assert {source['document_id'] for source in sources} == set(comparison_pair)
    assert all(source['page'] in [1, 2] and source['location_kind'] == 'page' for source in sources)
    assert 'Demo comparison' in result['token']['delta']
    conversation = client.get('/api/conversations/' + result['done']['conversation_id']).json()
    assert conversation['messages'][-1]['sources'] == sources


def test_ai_comparison_receives_tables_from_both_files(client, comparison_pair, monkeypatch):
    from app.api.routes import chat

    class FakeProvider:
        provider = 'test'

        async def stream(self, system, history, lang):
            if 'COVERAGE:' in system:
                assert 'contract-a.pdf' in system and 'contract-b.pdf' in system
                assert 'Wheat |  | 120' in system
                assert 'Page 2' in system
                assert 'not an exhaustive audit' in system
                yield 'Both documents price wheat at 120 [2] [4].'
            else:
                yield '{"title": "Contract comparison"}'

    monkeypatch.setattr(chat, 'get_provider', lambda _: FakeProvider())
    response = client.post('/api/chat', json={
        'mode': 'compare', 'message': 'Compare prices', 'document_ids': comparison_pair,
        'model': 'test', 'lang': 'en',
    })
    result = events(response)
    assert 'error' not in result
    assert 'price wheat' in result['token']['delta']
