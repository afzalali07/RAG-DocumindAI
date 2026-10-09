import asyncio
import json
import pytest
from openpyxl import Workbook
from app.core.spreadsheet import read_workbook, execute, Plan, answer
from app.llm.base import ChatMessage


@pytest.fixture
def emissions(tmp_path):
    w = Workbook()
    s = w.active
    s.title = 'Emission Data'
    s.append(['Factory', 'CO2 (t)', 'Emission Status'])
    for row in [('A', 10, 'High'), ('B', 20, 'Low'), ('C', 30, 'High')]:
        s.append(row)
    w.create_sheet('Summary').append(['Metric', 'Value'])
    path = tmp_path / 'emissions.xlsx'
    w.save(path)
    return path


class NoPlanner:
    async def stream(self, *args):
        raise AssertionError('An explicit categorical count should not require the model')
        yield ''


def test_emission_count_with_explicit_sheet(emissions):
    text, sheet, plan = asyncio.run(answer(
        emissions, 'What is the count of High in the Emission status in the Emission data ?',
        [], NoPlanner()))
    assert 'Row count: **2**' in text
    assert sheet['name'] == 'Emission Data'
    assert plan['filters'] == [{'column': 'Emission Status', 'operator': 'eq', 'value': 'High'}]
    assert plan['column'] == ''


def test_emission_sheet_followup_preserves_count(emissions):
    history = [
        ChatMessage(role='user', content='What is the count of High in the Emission status?'),
        ChatMessage(role='assistant', content='Please select a sheet: Emission Data, Summary'),
        ChatMessage(role='user', content='Emission Data'),
    ]
    text, _, _ = asyncio.run(answer(emissions, 'Emission Data', history, NoPlanner()))
    assert 'Row count: **2**' in text


def test_categorical_count_does_not_discard_extra_condition(emissions):
    from app.core.spreadsheet import direct_plan
    assert direct_plan(read_workbook(emissions),
                       'Count High in Emission Status in Emission Data where Factory is A') is None


def test_count_ambiguous_columns_still_requires_selection():
    from app.core.spreadsheet import direct_plan
    sheets = [{'name': name, 'headers': ['Status'], 'rows': [(2, ['High'])], 'missing': set()}
              for name in ('One', 'Two')]
    assert direct_plan(sheets, 'Count High in Status') is None

@pytest.fixture
def workbook(tmp_path):
    w=Workbook(); s=w.active; s.title='Marks'
    s.append(['Name','CO5 (20)','CO5 (100)','Team'])
    s.append(['Alice',16.5,82.5,'A'])
    s.append(['Bob',18,90,'B'])
    s.append(['Cara',18,90,'A'])
    s.append(['Dan',0,0,'B'])
    s.append(['Eve',None,None,'B'])
    p=tmp_path/'marks.xlsx'; w.save(p)
    return p

def test_exact_column_max_and_ties(workbook):
    text,_=execute(read_workbook(workbook),Plan(operation='max',column='CO5 (100)'))
    assert '**Max: 90**' in text
    assert 'Bob' in text and 'Cara' in text and 'Alice' not in text
    assert 'C3' in text and 'C4' in text

@pytest.mark.parametrize('operation,expected',[('sum','262.5'),('average','65.625'),('min','0')])
def test_aggregates(workbook,operation,expected):
    text,_=execute(read_workbook(workbook),Plan(operation=operation,column='CO5 (100)'))
    assert expected in text
    assert '4 numeric values' in text

def test_filter_count(workbook):
    plan=Plan(operation='count',filters=[{'column':'CO5 (100)','operator':'ge','value':85}])
    text,_=execute(read_workbook(workbook),plan)
    assert 'Row count: **2**' in text

def test_lookup_and_groups(workbook):
    text,_=execute(read_workbook(workbook),Plan(operation='rows',filters=[{'column':'Name','operator':'eq','value':'Alice'}]))
    assert 'Alice' in text and 'Bob' not in text
    text,_=execute(read_workbook(workbook),Plan(operation='group',group_by='Team',column='CO5 (100)',aggregate='average'))
    assert '86.25' in text and '45.0' in text

def test_ambiguous_column(workbook):
    with pytest.raises(ValueError,match='exact column'):
        execute(read_workbook(workbook),Plan(operation='max',column='CO5'))

def test_missing_formula_results_refused(workbook):
    from openpyxl import load_workbook
    w=load_workbook(workbook); w.active['C3']='=18*5'; w.save(workbook)
    with pytest.raises(ValueError,match='formulas without saved results'):
        execute(read_workbook(workbook),Plan(operation='max',column='CO5 (100)'))

def test_plan_cannot_run_code():
    with pytest.raises(ValueError): Plan.model_validate({'operation':'eval','code':'anything'})

def test_question_planning(workbook):
    class Fake:
        async def stream(self,system,messages,lang):
            assert 'CO5 (100)' in system
            yield json.dumps({'operation':'max','sheet':'Marks','column':'CO5 (100)'})
    text,_,_=asyncio.run(answer(workbook,'Who has highest CO5(100)?',[],Fake()))
    assert '**Max: 90**' in text
    text,_,_=asyncio.run(answer(workbook,'Who has highest CO5?',[],Fake()))
    assert 'Which column' in text


def test_wrong_column_plan_refused(workbook):
    class Wrong:
        async def stream(self, *args):
            yield json.dumps({'operation': 'max', 'column': 'CO5 (20)'})
    with pytest.raises(ValueError, match='different column'):
        asyncio.run(answer(workbook, 'Highest CO5(100)?', [], Wrong()))


def test_formula_only_trailing_row_is_not_lost(workbook):
    from openpyxl import load_workbook
    w = load_workbook(workbook)
    w.active['C7'] = '=100'
    w.save(workbook)
    with pytest.raises(ValueError, match='formulas without saved results'):
        execute(read_workbook(workbook), Plan(operation='max', column='CO5 (100)'))


@pytest.mark.parametrize('mode', ['rag', 'agent'])
def test_chat_executes_workbook_and_persists_sources(client, workbook, monkeypatch, mode):
    import shutil
    from app.api.routes import chat
    from app.config import get_settings
    from app.db.models import Document, Message
    from app.db.session import SessionLocal
    from sqlalchemy import select
    class Planner:
        async def stream(self, *args):
            yield json.dumps({'operation': 'max', 'column': 'CO5 (100)'})
    monkeypatch.setattr(chat, 'get_provider', lambda _: Planner())
    with SessionLocal() as db:
        document = Document(filename='marks.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', status='ready')
        db.add(document)
        db.commit()
        document_id = document.id
    shutil.copyfile(workbook, get_settings().upload_dir / f'{document_id}.xlsx')
    response = client.post('/api/chat', json={'message': 'Highest CO5(100)?', 'mode': mode, 'document_ids': [document_id]})
    assert response.status_code == 200
    assert '**Max: 90**' in response.text
    assert 'event: error' not in response.text
    if mode == 'agent': assert 'analyze_spreadsheet' in response.text
    with SessionLocal() as db:
        result = db.scalars(select(Message).where(Message.role == 'assistant').order_by(Message.created_at.desc())).first()
        assert '**Max: 90**' in result.content
        assert result.sources[0]['document_id'] == document_id
        assert result.sources[0]['label'] == 'Marks'


def test_screenshot_question_needs_no_model_plan(workbook):
    class BrokenPlanner:
        async def stream(self, *args):
            raise AssertionError('Clear ranking questions must not need a model plan')
            yield ''
    text, sheet, plan = asyncio.run(answer(workbook, 'which student has the highest mark in CO5(100) column', [], BrokenPlanner()))
    assert '**Max: 90**' in text
    assert sheet['name'] == 'Marks'
    assert plan['column'] == 'CO5 (100)'


def test_direct_plan_does_not_drop_filters(workbook):
    from app.core.spreadsheet import direct_plan
    assert direct_plan(read_workbook(workbook), 'which student has the highest mark in CO5(100) column where Team is A') is None


def test_invalid_plan_retried_with_selected_workbook(workbook):
    class Planner:
        calls = 0
        async def stream_json(self, system, messages, schema):
            self.calls += 1
            assert 'already selected' in system
            assert 'operation' in schema['properties']
            yield 'not json' if self.calls == 1 else json.dumps({'operation': 'average', 'column': 'CO5 (100)'})
    planner = Planner()
    text, _, _ = asyncio.run(answer(workbook, 'Average CO5(100)?', [], planner))
    assert '65.625' in text
    assert planner.calls == 2


def test_failed_plan_does_not_ask_for_document_selection(workbook):
    class Broken:
        async def stream(self, *args):
            yield 'invalid'
    with pytest.raises(ValueError, match='Your document is selected'):
        asyncio.run(answer(workbook, 'Average CO5(100)?', [], Broken()))


def test_ollama_sends_schema_and_preserves_gpu_setting(monkeypatch):
    import httpx
    from app.llm import providers
    captured = []
    def handler(request):
        captured.append(json.loads(request.content))
        return httpx.Response(200, text=json.dumps({'message': {'content': '{}'}})+'\n')
    original = httpx.AsyncClient
    monkeypatch.setattr(providers.httpx, 'AsyncClient', lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs))
    monkeypatch.setattr(providers.settings, 'ollama_num_gpu', 0)
    async def run():
        return [part async for part in providers.OllamaProvider().stream_json('plan', [], Plan.model_json_schema())]
    assert asyncio.run(run()) == ['{}']
    assert captured[0]['format'] == Plan.model_json_schema()
    assert captured[0]['options'] == {'temperature': 0, 'num_gpu': 0}
